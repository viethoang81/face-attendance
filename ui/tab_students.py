from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QTableWidget,
    QTableWidgetItem, QLineEdit, QComboBox, QPushButton,
    QLabel, QMessageBox, QHeaderView, QAbstractItemView
)
from PyQt5.QtCore import Qt, QTimer
from PyQt5.QtGui import QColor
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


STUDENT_STATUS_LABELS = {
    "DANG_HOC": "Đang học",
    "BAO_LUU": "Bảo lưu",
    "DINH_CHI": "Đình chỉ",
    "THOI_HOC": "Đã thôi học",
    "TOT_NGHIEP": "Tốt nghiệp",
}


class TabStudents(QWidget):
    def __init__(
        self,
        sv_svc,
        cls_svc,
        face_engine,
        current_user: dict | None = None,
    ):
        super().__init__()
        self.sv_svc = sv_svc
        self.cls_svc = cls_svc
        self.face_engine = face_engine
        self.current_user = dict(current_user or {})

        self._setup_ui()
        self._load_class_filter()
        self.load_data()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(8)

        search_row = QHBoxLayout()

        self.txt_search = QLineEdit()
        self.txt_search.setPlaceholderText(
            "Tìm kiếm theo MSSV hoặc họ tên..."
        )
        self.txt_search.setMinimumHeight(32)

        self._timer = QTimer()
        self._timer.setSingleShot(True)
        self._timer.timeout.connect(self._do_search)
        self.txt_search.textChanged.connect(
            lambda: self._timer.start(300)
        )

        self.cmb_lop = QComboBox()
        self.cmb_lop.setMinimumHeight(32)
        self.cmb_lop.currentIndexChanged.connect(self._do_search)

        self.cmb_status = QComboBox()
        self.cmb_status.setMinimumHeight(32)
        self.cmb_status.addItem("Tất cả trạng thái", "")
        self.cmb_status.addItem("Đang học", "DANG_HOC")
        self.cmb_status.addItem("Thôi học", "THOI_HOC")
        self.cmb_status.currentIndexChanged.connect(self._do_search)

        btn_refresh = QPushButton("Làm mới")
        btn_refresh.setMinimumHeight(32)
        btn_refresh.clicked.connect(self._refresh)

        search_row.addWidget(QLabel("Tìm:"))
        search_row.addWidget(self.txt_search, 3)
        search_row.addWidget(QLabel("Lớp:"))
        search_row.addWidget(self.cmb_lop, 1)
        search_row.addWidget(QLabel("Trạng thái:"))
        search_row.addWidget(self.cmb_status, 1)
        search_row.addWidget(btn_refresh)
        layout.addLayout(search_row)

        self.table = QTableWidget()
        self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels([
            "MSSV", "Họ và tên", "Lớp",
            "Ngày sinh", "Giới tính", "Email", "Số ĐT", "Trạng thái"
        ])
        self.table.horizontalHeader().setSectionResizeMode(
            1, QHeaderView.Stretch
        )
        self.table.horizontalHeader().setSectionResizeMode(
            5, QHeaderView.Stretch
        )
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.doubleClicked.connect(self._open_edit)
        layout.addWidget(self.table)

        self.lbl_status = QLabel("Tổng: 0 sinh viên")
        self.lbl_status.setStyleSheet("color: gray; font-size: 12px;")
        layout.addWidget(self.lbl_status)

        btn_row = QHBoxLayout()
        self._make_btn(btn_row, "+ Thêm mới", "#1D9E75", self._open_add)
        self.btn_edit = self._make_btn(
            btn_row, "Sửa / Xem", "#BA7517", self._open_edit
        )
        self.btn_remove = self._make_btn(
            btn_row,
            "Xóa / Cho thôi học",
            "#D85A30",
            self._delete,
        )
        self.btn_face = self._make_btn(
            btn_row,
            "Đăng ký khuôn mặt",
            "#534AB7",
            self._register_face,
        )
        btn_row.addStretch()
        self._make_btn(btn_row, "Xuất Excel", "#3B6D11", self._export)
        layout.addLayout(btn_row)

    def _make_btn(self, layout, text, color, handler):
        btn = QPushButton(text)
        btn.setMinimumHeight(34)
        btn.setStyleSheet(
            f"QPushButton {{ background:{color}; color:white; "
            f"border-radius:6px; padding:0 14px; font-weight:500; }}"
            f"QPushButton:hover {{ opacity:0.85; }}"
        )
        btn.clicked.connect(handler)
        layout.addWidget(btn)
        return btn

    def _load_class_filter(self):
        self.cmb_lop.clear()
        self.cmb_lop.addItem("Tất cả lớp", "")

        for classroom in self.cls_svc.get_all():
            self.cmb_lop.addItem(
                classroom["tenLop"],
                classroom["maLop"],
            )

    @staticmethod
    def _is_student_complete(
        sv: dict,
        registered_set: set[str],
    ) -> tuple[bool, str]:
        """Kiểm tra sinh viên đủ thông tin và có mẫu khuôn mặt."""

        missing = []
        required_fields = {
            "mssv": "MSSV",
            "hoTen": "Họ tên",
            "maLop": "Lớp",
            "ngaySinh": "Ngày sinh",
            "email": "Email",
            "soDienThoai": "Số điện thoại",
        }

        for field, label in required_fields.items():
            if not sv.get(field, "").strip():
                missing.append(label)

        if sv.get("mssv", "").strip() not in registered_set:
            missing.append("Khuôn mặt")

        if missing:
            return False, "Chưa hoàn tất - Thiếu: " + ", ".join(missing)

        return True, "Đã hoàn tất thông tin và đăng ký khuôn mặt"

    def load_data(self, rows=None):
        data = rows if rows is not None else self.sv_svc.get_all()
        registered_set = self.face_engine.get_registered_mssv_set()

        self.table.setRowCount(0)

        for sv in data:
            row = self.table.rowCount()
            self.table.insertRow(row)

            is_complete, tooltip = self._is_student_complete(
                sv,
                registered_set,
            )
            status = sv.get("trangThai", "")
            if status == "THOI_HOC":
                background = QColor("#FCEBEB")
                tooltip = "Đã thôi học - chỉ được xem thông tin"
            else:
                background = (
                    QColor("#E1F5EE") if is_complete else QColor("#FFFFFF")
                )

            for col, key in enumerate([
                "mssv",
                "hoTen",
                "tenLop",
                "ngaySinh",
                "gioiTinh",
                "email",
                "soDienThoai",
                "trangThai",
            ]):
                value = sv.get(key, "")
                if key == "trangThai":
                    value = STUDENT_STATUS_LABELS.get(value, value)
                item = QTableWidgetItem(str(value))
                item.setTextAlignment(Qt.AlignCenter)
                item.setBackground(background)
                item.setToolTip(tooltip)
                self.table.setItem(row, col, item)

        self.lbl_status.setText(f"Tổng: {len(data)} sinh viên")

    def _do_search(self):
        keyword = self.txt_search.text().strip()
        ma_lop = self.cmb_lop.currentData()
        status = self.cmb_status.currentData() or ""
        self.load_data(self.sv_svc.search(keyword, ma_lop, status))

    def _refresh(self):
        self.txt_search.clear()
        self.cmb_lop.setCurrentIndex(0)
        self.cmb_status.setCurrentIndex(0)
        self._load_class_filter()
        self.load_data()

    def _selected_mssv(self) -> str | None:
        row = self.table.currentRow()
        if row < 0:
            QMessageBox.warning(
                self,
                "Chú ý",
                "Vui lòng chọn một sinh viên.",
            )
            return None

        return self.table.item(row, 0).text()

    def _open_add(self):
        from ui.dialogs import StudentDialog

        dlg = StudentDialog(self.sv_svc, self.cls_svc, parent=self)
        if dlg.exec_():
            self._refresh()

    def _open_edit(self):
        mssv = self._selected_mssv()
        if not mssv:
            return

        student = self.sv_svc.get_by_mssv(mssv)
        if not student:
            return

        from ui.dialogs import StudentDialog

        dlg = StudentDialog(
            self.sv_svc,
            self.cls_svc,
            existing=student,
            read_only=student.get("trangThai") != "DANG_HOC",
            parent=self,
        )
        if dlg.exec_() and student.get("trangThai") == "DANG_HOC":
            self.face_engine.load_encodings()
            self._refresh()

    def _delete(self):
        mssv = self._selected_mssv()
        if not mssv:
            return

        can_purge, reason = self.sv_svc.can_hard_delete(
            mssv,
            actor=self.current_user,
        )
        if can_purge:
            reply = QMessageBox.warning(
                self,
                "Xóa vĩnh viễn sinh viên",
                "Sinh viên chưa phát sinh dữ liệu học tập. Hồ sơ và mẫu "
                "khuôn mặt sẽ bị xóa vĩnh viễn và không thể khôi phục.",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No,
            )
            if reply != QMessageBox.Yes:
                return
            ok, message = self.sv_svc.purge_student(
                mssv,
                actor=self.current_user,
            )
        else:
            if "quyền" in reason.lower() or "không tìm thấy" in reason.lower():
                QMessageBox.warning(self, "Không thể thao tác", reason)
                return
            student = self.sv_svc.get_by_mssv(mssv)
            if student and student.get("trangThai") == "THOI_HOC":
                QMessageBox.information(
                    self,
                    "Sinh viên đã thôi học",
                    "Sinh viên đã ở trạng thái thôi học và vẫn có dữ liệu "
                    "học vụ nên không thể xóa vĩnh viễn.",
                )
                return
            from ui.dialogs import StudentWithdrawalDialog

            dialog = StudentWithdrawalDialog(mssv, reason, self)
            if not dialog.exec_():
                return
            ok, message = self.sv_svc.withdraw_student(
                mssv,
                actor=self.current_user,
                effective_date=dialog.effective_date(),
            )
        if ok:
            self.face_engine.load_encodings()

        QMessageBox.information(self, "Kết quả", message)

        if ok:
            self._refresh()

    def _register_face(self):
        mssv = self._selected_mssv()
        if not mssv:
            return

        student = self.sv_svc.get_by_mssv(mssv)
        if not student:
            return
        if student.get("trangThai") != "DANG_HOC":
            QMessageBox.warning(
                self,
                "Không thể đăng ký khuôn mặt",
                "Sinh viên không ở trạng thái Đang học; chỉ được xem "
                "thông tin.",
            )
            return

        registered_set = self.face_engine.get_registered_mssv_set()
        if mssv in registered_set:
            reply = QMessageBox.question(
                self,
                "Khuôn mặt đã tồn tại",
                f"Sinh viên '{mssv}' đã được đăng ký khuôn mặt trước đó.\n\n"
                "Bạn có muốn đăng ký lại không?\n\n"
                "Mẫu cũ chỉ bị thay thế sau khi bạn chụp thành công "
                "đủ 10 mẫu mới.",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No,
            )
            if reply != QMessageBox.Yes:
                return

        from ui.dialogs import FaceRegisterDialog

        dlg = FaceRegisterDialog(mssv, self.face_engine, parent=self)
        if dlg.exec_():
            self._refresh()

    def _export(self):
        from PyQt5.QtWidgets import QFileDialog
        import pandas as pd

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Lưu file Excel",
            "danh_sach_sinh_vien.xlsx",
            "Excel (*.xlsx)",
        )
        if not path:
            return

        columns = [
            ("mssv", "MSSV"),
            ("hoTen", "Họ tên"),
            ("tenLop", "Lớp"),
            ("ngaySinh", "Ngày sinh"),
            ("gioiTinh", "Giới tính"),
            ("email", "Email"),
            ("soDienThoai", "Số ĐT"),
            ("trangThai", "Trạng thái"),
        ]

        try:
            data = self.sv_svc.get_all()
            source_keys = [key for key, _ in columns]
            headers = [label for _, label in columns]

            dataframe = pd.DataFrame(data).reindex(columns=source_keys)
            dataframe.columns = headers
            dataframe["Trạng thái"] = dataframe["Trạng thái"].map(
                lambda value: STUDENT_STATUS_LABELS.get(value, value)
            )

            # Giữ MSSV và số điện thoại dạng chữ để không mất số 0 đầu.
            for column_name in ("MSSV", "Số ĐT"):
                dataframe[column_name] = dataframe[column_name].map(
                    lambda value: "" if pd.isna(value) else str(value)
                )

            sheet_name = "Danh sách sinh viên"

            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                dataframe.to_excel(
                    writer,
                    index=False,
                    sheet_name=sheet_name,
                )
                worksheet = writer.sheets[sheet_name]

                # Đồng bộ với file Lịch sử điểm danh và Báo cáo chuyên cần.
                header_fill = PatternFill(
                    start_color="1D9E75",
                    end_color="1D9E75",
                    fill_type="solid",
                )

                for cell in worksheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")

                # Cùng cơ chế tự tính độ rộng cột như các báo cáo còn lại.
                for column_index, column_cells in enumerate(
                    worksheet.columns,
                    start=1,
                ):
                    max_length = max(
                        len(str(cell.value or ""))
                        for cell in column_cells
                    )
                    worksheet.column_dimensions[
                        get_column_letter(column_index)
                    ].width = min(max(max_length + 4, 10), 40)

            QMessageBox.information(
                self,
                "Thành công",
                f"Đã xuất {len(data)} sinh viên ra:\n{path}",
            )

        except Exception as exc:
            QMessageBox.critical(
                self,
                "Xuất Excel thất bại",
                f"Không thể tạo file Excel:\n{exc}",
            )
