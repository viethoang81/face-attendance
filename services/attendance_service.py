import logging
from core.database import Database
from core import attendance_status as A

logger = logging.getLogger(__name__)


class AttendanceService:
    def __init__(self, db: Database):
        self.db = db

    def get_history(self, maLop="", mssv="",
                    ngay_from="", ngay_to="") -> list[dict]:
        sql = """
            SELECT dd.id,
                   sv.mssv,
                   sv.hoTen,
                   lh.tenLop,
                   COALESCE(mh.tenMon, '') AS tenMon,
                   ca.tenCa,
                   ca.ngayHoc,
                   COALESCE(dd.thoiGian, '') AS thoiGian,
                   CASE
                        WHEN dd.id IS NULL THEN 'ABSENT'
                        ELSE COALESCE(dd.trangThai, 'PRESENT')
                   END AS trangThai
            FROM CaHoc ca
            JOIN LopHoc lh ON ca.maLop = lh.maLop
            JOIN SinhVien sv ON sv.maLop = ca.maLop
            LEFT JOIN MonHoc mh ON ca.maMon = mh.maMon
            LEFT JOIN DiemDanh dd
                   ON dd.caHocId = ca.id
                  AND dd.mssv = sv.mssv
            WHERE 1=1
                AND (
                    COALESCE(ca.daMoDiemDanh, 0) = 1
                    OR EXISTS (
                        SELECT 1
                        FROM DiemDanh dd2
                        WHERE dd2.caHocId = ca.id
                    )
                )    
        """
        params = []
        if maLop:
            sql += " AND ca.maLop=?"
            params.append(maLop)
        if mssv:
            sql += " AND sv.mssv LIKE ?"
            params.append(f"%{mssv}%")
        if ngay_from:
            sql += " AND ca.ngayHoc>=?"
            params.append(ngay_from)
        if ngay_to:
            sql += " AND ca.ngayHoc<=?"
            params.append(ngay_to)
        sql += " ORDER BY ca.ngayHoc DESC, ca.gioBatDau DESC, sv.hoTen"
        with self.db.conn() as c:
            rows = [dict(r) for r in c.execute(sql, params)]
        for row in rows:
            row["thoiGian"] = self._time_only(row.get("thoiGian"))
        return rows

    def get_session_summary(self, ca_hoc_id: int) -> list[dict]:
        """Danh sách sinh viên đã điểm danh trong 1 ca."""
        with self.db.conn() as c:
            rows = c.execute("""
                SELECT dd.mssv, sv.hoTen,
                       dd.thoiGian, dd.trangThai
                FROM DiemDanh dd
                JOIN SinhVien sv ON dd.mssv = sv.mssv
                WHERE dd.caHocId = ?
                  AND dd.isVoided = 0
                ORDER BY dd.thoiGian
            """, (ca_hoc_id,)).fetchall()
            return [dict(r) for r in rows]

    def get_session_full_report(self, ca_hoc_id: int,
                                ma_lop: str) -> list[dict]:
        """
        Báo cáo đầy đủ 1 ca: có mặt + vắng mặt.
        Lấy toàn bộ SV trong lớp, ai không điểm danh → Vắng.
        """
        with self.db.conn() as c:
            ca_info = c.execute("""
                SELECT ca.tenCa, ca.ngayHoc, ca.gioBatDau, ca.gioKetThuc,
                       lh.tenLop,
                       COALESCE(mh.tenMon, '') AS tenMon
                FROM CaHoc ca
                JOIN LopHoc lh ON ca.maLop = lh.maLop
                LEFT JOIN MonHoc mh ON ca.maMon = mh.maMon
                WHERE ca.id=?
            """, (ca_hoc_id,)).fetchone()

            # Lấy tất cả SV trong lớp
            all_sv = c.execute("""
                SELECT mssv, hoTen FROM SinhVien
                WHERE maLop=? ORDER BY hoTen
            """, (ma_lop,)).fetchall()

            # Lấy danh sách đã điểm danh
            attended = {
                r["mssv"]: r for r in [
                    dict(x) for x in c.execute("""
                        SELECT mssv, thoiGian, trangThai
                        FROM DiemDanh WHERE caHocId=?
                    """, (ca_hoc_id,)).fetchall()
                ]
            }

        ca_info = dict(ca_info) if ca_info else {}
        gio_hoc = ""
        if ca_info.get("gioBatDau") and ca_info.get("gioKetThuc"):
            gio_hoc = f"{ca_info['gioBatDau']} - {ca_info['gioKetThuc']}"

        report = []
        for sv in all_sv:
            mssv   = sv["mssv"]
            ho_ten = sv["hoTen"]
            if mssv in attended:
                report.append({
                    "mssv":      mssv,
                    "hoTen":     ho_ten,
                    "tenLop":    ca_info.get("tenLop", ""),
                    "tenMon":    ca_info.get("tenMon", ""),
                    "tenCa":     ca_info.get("tenCa", ""),
                    "ngayHoc":   ca_info.get("ngayHoc", ""),
                    "gioHoc":    gio_hoc,
                    "thoiGian":  self._time_only(attended[mssv]["thoiGian"]),
                    "trangThai": A.normalize(attended[mssv]["trangThai"])
                })
            else:
                report.append({
                    "mssv":      mssv,
                    "hoTen":     ho_ten,
                    "tenLop":    ca_info.get("tenLop", ""),
                    "tenMon":    ca_info.get("tenMon", ""),
                    "tenCa":     ca_info.get("tenCa", ""),
                    "ngayHoc":   ca_info.get("ngayHoc", ""),
                    "gioHoc":    gio_hoc,
                    "thoiGian":  "",
                    "trangThai": A.ABSENT
                })
        return report

    # ── Export lịch sử đầy đủ (nhiều ca) ──────────
    def export_excel_history(self, records: list[dict],
                             path: str) -> bool:
        """Xuất lịch sử điểm danh nhiều ca."""
        try:
            import pandas as pd
            df = pd.DataFrame(records)
            if "trangThai" in df.columns:
                df["trangThai"] = df["trangThai"].map(A.label)
            # Chọn cột ổn định; các khóa kỹ thuật chỉ dùng nội bộ.
            df = df.reindex(columns=[
                "mssv", "hoTen", "maLopHocPhan", "tenLop", "tenMon",
                "tenCa", "ngayHoc", "gioHoc", "thoiGian", "trangThai"
            ], fill_value="")
            df.columns = [
                "MSSV", "Họ tên", "Mã lớp học phần", "Lớp học phần",
                "Môn học", "Ca học", "Ngày học", "Giờ học",
                "Thời gian điểm danh", "Trạng thái"
            ]
            self._save_excel(df, path, "Lich su diem danh")
            return True
        except Exception:
            logger.exception("[AttendanceService] Cannot export attendance history")
            return False

    # ── Export báo cáo 1 ca (có mặt + vắng) ───────
    def export_excel_session(self, records: list[dict],
                             path: str,
                             ten_ca: str = "") -> bool:
        """Xuất báo cáo 1 ca học."""
        try:
            import pandas as pd
            df = pd.DataFrame(records)
            if "trangThai" in df.columns:
                df["trangThai"] = df["trangThai"].map(A.label)
            df = df.reindex(columns=[
                "mssv", "hoTen", "tenLop", "tenMon", "tenCa",
                "ngayHoc", "gioHoc", "thoiGian", "trangThai"
            ], fill_value="")
            df.columns = [
                "MSSV", "Họ tên", "Lớp", "Môn học", "Ca học",
                "Ngày học", "Giờ học", "Thời gian điểm danh",
                "Trạng thái"
            ]
            # Tô màu theo trạng thái
            sheet = ten_ca[:30] if ten_ca else "Bao cao ca hoc"
            self._save_excel(df, path, sheet)
            return True
        except Exception:
            logger.exception("[AttendanceService] Cannot export session report")
            return False

    # ── Hàm export dùng chung ──────────────────────
    @staticmethod
    def _save_excel(df, path: str, sheet_name: str):
        """Lưu DataFrame ra Excel với định dạng đẹp."""
        import pandas as pd
        from openpyxl.styles import PatternFill, Font, Alignment
        from openpyxl.utils import get_column_letter

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False,
                        sheet_name=sheet_name)
            ws = writer.sheets[sheet_name]

            # Header đậm + nền xanh
            header_fill = PatternFill(
                start_color="1D9E75", end_color="1D9E75",
                fill_type="solid"
            )
            for cell in ws[1]:
                cell.font      = Font(bold=True, color="FFFFFF")
                cell.fill      = header_fill
                cell.alignment = Alignment(horizontal="center")

            # Tô màu hàng vắng mặt
            red_fill = PatternFill(
                start_color="FCEBEB", end_color="FCEBEB",
                fill_type="solid"
            )
            for row in ws.iter_rows(min_row=2):
                last_cell = row[-1]
                if last_cell.value in A.ABSENCE_LABELS or \
                        last_cell.value in ("Vang mat", "Vắng mặt"):
                    for cell in row:
                        cell.fill = red_fill

            # Các mã định danh phải là văn bản. Nếu để General, một số
            # trình đọc Excel sẽ tự đổi "001234" thành số 1234 và làm mất
            # số 0 đầu dù dữ liệu nguồn trong SQLite là TEXT.
            text_headers = {"MSSV", "Mã lớp học phần"}
            for column_index, header_cell in enumerate(ws[1], start=1):
                if header_cell.value not in text_headers:
                    continue
                for row_index in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_index, column=column_index)
                    cell.value = "" if cell.value is None else str(cell.value)
                    cell.number_format = "@"
                    cell.quotePrefix = True

            # Tự động độ rộng cột
            for col_idx, col in enumerate(ws.columns, 1):
                max_len = max(
                    len(str(c.value or "")) for c in col
                )
                header = str(col[0].value or "")
                min_width = {
                    "Ngày học": 14,
                    "Giờ học": 16,
                    "Thời gian": 22,
                    "Thời gian điểm danh": 22,
                }.get(header, 0)
                ws.column_dimensions[
                    get_column_letter(col_idx)
                ].width = min(max(max_len + 4, min_width), 40)

    # ── Giữ backward compatible ────────────────────
    def export_excel(self, records: list[dict],
                     path: str) -> bool:
        """Backward compatible — tự detect loại records."""
        if not records:
            return False
        if "caHocId" in records[0] or "maLopHocPhan" in records[0]:
            return self.export_excel_history(records, path)
        if "gioHoc" in records[0]:
            return self.export_excel_session(records, path)
        if "tenCa" in records[0]:
            return self.export_excel_history(records, path)
        return self.export_excel_session(records, path)
    
    @staticmethod
    def _time_only(value) -> str:
        text = str(value or "").strip()
        if not text:
            return ""
        if " " in text:
            return text.split()[-1][:8]
        if "T" in text:
            return text.split("T")[-1][:8]
        return text[:8]
    
    # Báo cáo chuyên cần
    def get_attendance_report(self, maLop="", maMon="",
                              ngay_from="", ngay_to="") -> dict:
        if not maLop:
            return {"threshold": A.DEFAULT_MIN_RATE, "total_sessions": 0,
                    "sessions_held": 0, "students": []}

        with self.db.conn() as c:
            threshold = A.DEFAULT_MIN_RATE
            so_ca_ke_hoach = 0
            if maMon:
                r = c.execute(
                    "SELECT soTiet, tiLeToiThieu FROM MonHoc WHERE maMon=?",
                    (maMon,),
                ).fetchone()
                if r:
                    if r["tiLeToiThieu"]:
                        threshold = int(r["tiLeToiThieu"])
                    if r["soTiet"]:
                        so_ca_ke_hoach = round(int(r["soTiet"]) / A.SO_TIET_MOI_CA)

            sql = (
                "SELECT id FROM CaHoc ca WHERE maLop=? AND ("
                "COALESCE(daMoDiemDanh,0)=1 "
                "OR EXISTS(SELECT 1 FROM DiemDanh d WHERE d.caHocId=ca.id))"
            )
            params = [maLop]
            if maMon:
                sql += " AND ca.maMon=?"
                params.append(maMon)
            if ngay_from:
                sql += " AND ca.ngayHoc>=?"
                params.append(ngay_from)
            if ngay_to:
                sql += " AND ca.ngayHoc<=?"
                params.append(ngay_to)
            session_ids = [row["id"] for row in c.execute(sql, params)]
            sessions_held = len(session_ids)

            total = so_ca_ke_hoach if (maMon and so_ca_ke_hoach > 0) else sessions_held

            students = [
                dict(r) for r in c.execute(
                    "SELECT mssv, hoTen FROM SinhVien WHERE maLop=? ORDER BY hoTen",
                    (maLop,),
                )
            ]

            present_count: dict[str, int] = {}
            if session_ids and students:
                placeholders = ",".join("?" * len(session_ids))
                rows = c.execute(
                    f"SELECT mssv, trangThai FROM DiemDanh "
                    f"WHERE caHocId IN ({placeholders})",
                    session_ids,
                ).fetchall()
                for row in rows:
                    if A.is_present(row["trangThai"]):
                        present_count[row["mssv"]] = \
                            present_count.get(row["mssv"], 0) + 1

            result = []
            for sv in students:
                co_mat = present_count.get(sv["mssv"], 0)
                co_mat_hieu_luc = min(co_mat, total) if total else co_mat
                ty_le = round(co_mat_hieu_luc / total * 100, 1) if total else 0.0
                so_vang = max(sessions_held - co_mat, 0)   # chỉ tính buổi ĐÃ học

                buoi_con_lai  = max(total - sessions_held, 0)
                co_mat_toi_da = min(co_mat + buoi_con_lai, total) if total else co_mat
                ty_le_toi_da  = round(co_mat_toi_da / total * 100, 1) if total else 0.0

                if total <= 0:
                    trang_thai_thi = ""
                elif ty_le >= threshold:
                    trang_thai_thi = "ĐỦ ĐIỀU KIỆN"
                elif ty_le_toi_da < threshold:
                    trang_thai_thi = "CẤM THI"
                else:
                    trang_thai_thi = ""

                result.append({
                    "mssv": sv["mssv"],
                    "hoTen": sv["hoTen"],
                    "tong_buoi": total,
                    "so_buoi_da_hoc": sessions_held,
                    "so_co_mat": co_mat,
                    "so_vang": so_vang,
                    "ty_le": ty_le,
                    "trang_thai_thi": trang_thai_thi,
                    "du_dieu_kien": (trang_thai_thi == "ĐỦ ĐIỀU KIỆN"), 
                })
            result.sort(key=lambda x: x["ty_le"])  # chuyên cần kém lên đầu

            return {
                "threshold": threshold,
                "total_sessions": total,
                "sessions_held": sessions_held,
                "students": result,
            }

    def export_excel_report(self, records: list[dict],
                            path: str, threshold: int = 75) -> bool:
        """Xuất báo cáo chuyên cần; tô đỏ sinh viên bị cấm thi."""
        try:
            import pandas as pd
            from openpyxl.styles import PatternFill, Font, Alignment
            from openpyxl.utils import get_column_letter

            df = pd.DataFrame([{
                "MSSV": r["mssv"], "Họ tên": r["hoTen"],
                "Số buổi": r["tong_buoi"], "Có mặt": r["so_co_mat"],
                "Vắng": r["so_vang"], "Tỷ lệ (%)": r["ty_le"],
                "Điều kiện dự thi":  r.get("trang_thai_thi", ""),
            } for r in records])

            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                sheet = f"Chuyen can (>={threshold}%)"[:31]
                df.to_excel(writer, index=False, sheet_name=sheet)
                ws = writer.sheets[sheet]
                header_fill = PatternFill("solid", fgColor="1D9E75")
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                green = PatternFill("solid", fgColor="E8F6EF")
                red   = PatternFill("solid", fgColor="FCEBEB")
                for row in ws.iter_rows(min_row=2):
                    status = row[-1].value    
                    if status == "CẤM THI":
                        for cell in row:
                            cell.fill = red
                    elif status == "ĐỦ ĐIỀU KIỆN":
                        for cell in row:
                            cell.fill = green
                for idx, col in enumerate(ws.columns, 1):
                    width = max(len(str(x.value or "")) for x in col) + 4
                    ws.column_dimensions[get_column_letter(idx)].width = \
                        min(max(width, 10), 40)
            return True
        except Exception:
            logger.exception("[AttendanceService] Cannot export attendance report")
            return False
